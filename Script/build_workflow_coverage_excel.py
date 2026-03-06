import argparse
import glob
import json
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

CATEGORIES = [
    "FCMRInspectionMode",
    "AVViewerApplication",
    "CcaApplication",
    "AvaApplication",
    "FunctionalCardiacCTApplication",
]


def parse_percentage(pct_str: str) -> float:
    if not pct_str:
        return 0.0
    s = str(pct_str).strip().replace("%", "")
    try:
        return float(s) / 100.0
    except Exception:
        return 0.0


def load_daily_data(logs_dir: str):
    daily = {}
    pattern = os.path.join(logs_dir, "data*.json")
    for path in glob.glob(pattern):
        base = os.path.basename(path)
        date_token = base.replace("data", "").replace(".json", "")
        try:
            d = datetime.strptime(date_token, "%Y-%m-%d").date()
        except Exception:
            continue
        try:
            with open(path, "r", encoding="utf-8") as f:
                top = json.load(f)
        except Exception:
            continue
        row = {}
        for cat in CATEGORIES:
            val = 0.0
            if cat in top:
                try:
                    inner = json.loads(top[cat])
                    pct = inner.get("Percentage :", "0%")
                    val = parse_percentage(pct)
                except Exception:
                    val = 0.0
            row[cat] = val
        daily[d] = row
    return daily


def build_workbook(daily_map: dict, output_path: str):
    dates = sorted(daily_map.keys())
    wb = Workbook()
    ws = wb.active
    ws.title = "Coverage"

    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.cell(row=1, column=1, value="Application").font = header_font
    ws.cell(row=1, column=1).fill = header_fill
    ws.cell(row=1, column=1).alignment = center
    ws.cell(row=1, column=1).border = border

    for idx, d in enumerate(dates, start=2):
        c = ws.cell(row=1, column=idx, value=d.strftime("%d-%b"))
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border

    for r, cat in enumerate(CATEGORIES, start=2):
        c = ws.cell(row=r, column=1, value=cat)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = border
        for idx, d in enumerate(dates, start=2):
            val = daily_map.get(d, {}).get(cat, 0.0)
            cell = ws.cell(row=r, column=idx, value=val)
            cell.number_format = "0%"
            cell.alignment = center
            cell.border = border

    ws.column_dimensions[get_column_letter(1)].width = 34
    for col in range(2, len(dates) + 2):
        ws.column_dimensions[get_column_letter(col)].width = 10

    ws.freeze_panes = "B2"

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="Build Workflow Coverage Excel from data*.json files")
    parser.add_argument("--logs", required=True, help="Path to folder containing data*.json files")
    parser.add_argument("--out", required=True, help="Output Excel file path")
    args = parser.parse_args()

    daily = load_daily_data(args.logs)
    if not daily:
        build_workbook({}, args.out)
        return

    build_workbook(daily, args.out)


if __name__ == "__main__":
    main()
